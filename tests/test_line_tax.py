"""Thuế dòng: ưu tiên số khai báo, phép tính theo dữ liệu và chỉ đọc Excel cũ."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
import parser
import excel_manager as excel
from invoice_tax import calculate_line_tax
from test_startup_performance import SAMPLE_XML


class LineTaxTests(unittest.TestCase):
    def invoice(self, rate='10%', extra='', amount='1000'):
        xml = SAMPLE_XML.replace('<TSuat>10%</TSuat>', f'<TSuat>{rate}</TSuat>{extra}').replace('<ThTien>1000</ThTien>',f'<ThTien>{amount}</ThTien>')
        return parser.parse_xml_invoice(xml)

    def test_declared_tax_preserved_even_zero_and_negative(self):
        for tag in ('TThue', 'VATAmount', 'TaxAmount'):
            for value in (25, 0, -25):
                item = self.invoice(extra=f'<{tag}>{value}</{tag}>')['hang_hoa'][0]
                self.assertEqual(item['tien_thue'],value)
                self.assertEqual(item['tong_tien_dong'],1000+value)

    def test_extension_tax_amount_not_tax_rate(self):
        extra = '<TTKhac><TTin><TTruong>Thuế suất</TTruong><DLieu>8</DLieu></TTin><TTin><TTruong>Tiền thuế GTGT</TTruong><DLieu>25</DLieu></TTin></TTKhac>'
        self.assertEqual(self.invoice(extra=extra)['hang_hoa'][0]['tien_thue'],25)
        self.assertEqual(self.invoice(extra=extra+'<TThue>20</TThue>')['hang_hoa'][0]['tien_thue'],20)

    def test_decimal_zero_exempt_unknown_and_negative_adjustment(self):
        for rate, expected in [('8%',80),('10%',100),('1,5%',15),('1.5%',15),('0%',0),('KCT',0),('KKKNT',0),('',None),('KHAC',None)]:
            item=self.invoice(rate=rate)['hang_hoa'][0]
            self.assertEqual(item['tien_thue'],expected,rate)
            self.assertEqual(item['tong_tien_dong'],None if expected is None else 1000+expected)
        self.assertEqual(self.invoice(amount='-1000')['hang_hoa'][0]['tien_thue'],-100)
        self.assertEqual(calculate_line_tax(1.05,'10%'),0.11)
        self.assertEqual(calculate_line_tax(1000,0),0)

    def test_missing_workbook_tax_can_be_calculated_without_modifying_explicit_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'old.xlsx'
            excel.init_blank_excel_file(str(path))
            excel.save_invoices_to_excel([self.invoice()],str(path))
            wb=openpyxl.load_workbook(path)
            ws=wb['ChiTietHangHoa']
            ws.cell(2,17).value=None
            ws.cell(2,18).value=None
            wb.save(path);wb.close()
            before=path.read_bytes()
            item=excel.read_excel_invoice_details(str(path))['invoices'][0]['hang_hoa'][0]
            self.assertEqual(item['tien_thue'],100)
            self.assertEqual(item['tong_tien_dong'],1100)
            self.assertEqual(path.read_bytes(),before)
            wb=openpyxl.load_workbook(path);wb['ChiTietHangHoa'].cell(2,17).value=0;wb.save(path);wb.close()
            before=path.read_bytes()
            self.assertEqual(excel.read_excel_invoice_details(str(path))['invoices'][0]['hang_hoa'][0]['tien_thue'],0)
            self.assertEqual(path.read_bytes(),before)

    def pdf(self, lines, rate='', vat=100, base=1000):
        text = f'HÓA ĐƠN GIÁ TRỊ GIA TĂNG\nKý hiệu: C26TAA\nSố hóa đơn: 123\nNgười bán: Nhà cung cấp mẫu\nMã số thuế: 0123456789\nNgười mua: Khách hàng mẫu\nMã số thuế: 9876543210\n{rate}\n{lines}\nCộng tiền hàng: {base}\nTiền thuế GTGT: {vat}\nTổng tiền thanh toán: {base+vat}'
        with patch.object(parser.pypdf,'PdfReader') as reader, patch.object(parser,'extract_layout_text_from_pdf',return_value=text):
            reader.return_value.attachments={}
            result=parser.parse_pdf_invoice(BytesIO(b'fixture'),filename='fixture.pdf')
        self.assertTrue(result['success'],result)
        return result

    def test_pdf_line_or_verified_header_rate_and_single_item_tax(self):
        cases=[('1 Sản phẩm mẫu Cái 1 1000 1000 10%','',100,'10%'),
               ('1 Sản phẩm mẫu Cái 1 1000 1000','Thuế suất GTGT: 10%',100,'10%'),
               ('1 Sản phẩm mẫu Cái 1 1000 1000','',100,'')]
        for lines,rate,expected,label in cases:
            item=self.pdf(lines,rate)['hang_hoa'][0]
            self.assertEqual(item['tien_thue'],expected)
            self.assertEqual(item['tong_tien_dong'],1100)
            self.assertEqual(item['thue_suat'],label)

    def test_pdf_multiple_items_without_rates_are_not_assumed_eight_percent(self):
        items=self.pdf('1 Sản phẩm mẫu Cái 1 400 400\n2 Dịch vụ mẫu Gói 1 600 600')['hang_hoa']
        self.assertEqual(len(items),2)
        self.assertTrue(all(item['tien_thue'] is None and item['thue_suat']=='' for item in items))
        items=self.pdf('1 Sản phẩm mẫu Cái 1 400 400\n2 Dịch vụ mẫu Gói 1 600 600','Thuế suất GTGT: 10%')['hang_hoa']
        self.assertEqual([item['tien_thue'] for item in items],[40,60])
