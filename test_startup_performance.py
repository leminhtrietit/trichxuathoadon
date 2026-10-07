"""Kiểm tra trong thư mục tạm, không thay đổi dữ liệu/cấu hình người dùng."""
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import excel_manager as excel
from parser import parse_xml_invoice

SAMPLE_XML = '''<HDon><DLHDon><TTChung><KHMSHDon>1</KHMSHDon>
<KHHDon>C26TAA</KHHDon><SHDon>123</SHDon><NLap>2026-10-07</NLap></TTChung>
<NDHDon><NBan><Ten>Nhà cung cấp mẫu</Ten><MST>0123456789</MST></NBan>
<NMua><Ten>Khách hàng mẫu</Ten><MST>9876543210</MST></NMua>
<DSHHDVu><HHDVu><THHDVu>Sản phẩm mẫu</THHDVu><SLuong>1</SLuong><DGia>1000</DGia>
<ThTien>1000</ThTien><TSuat>10%</TSuat></HHDVu></DSHHDVu>
<TToan><TgTCThue>1000</TgTCThue><TgTThue>100</TgTThue><TgTTTBSo>1100</TgTTTBSo></TToan>
</NDHDon></DLHDon></HDon>'''


class StartupPerformanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = str(Path(self.temp.name) / 'invoices.xlsx')
        self.invoice = parse_xml_invoice(SAMPLE_XML, filename='sample.xml')
        excel._summary_cache.clear()
        self.assertTrue(excel.init_blank_excel_file(self.path)['success'])

    def test_cache_reuses_read_and_does_not_share_mutable_rows(self):
        excel.save_invoices_to_excel([self.invoice], self.path)
        with patch.object(excel.openpyxl, 'load_workbook', wraps=openpyxl.load_workbook) as load:
            first = excel.read_excel_summary(self.path)
            first['rows'][0]['nb_ten'] = 'changed by caller'
            second = excel.read_excel_summary(self.path)
            self.assertEqual(load.call_count, 1)
            self.assertEqual(second['rows'][0]['nb_ten'], 'Nhà cung cấp mẫu')

    def test_parallel_reads_load_workbook_once(self):
        with patch.object(excel.openpyxl, 'load_workbook', wraps=openpyxl.load_workbook) as load:
            with ThreadPoolExecutor(max_workers=4) as pool:
                results = list(pool.map(excel.read_excel_summary, [self.path] * 4))
            self.assertEqual(load.call_count, 1)
            self.assertTrue(all(r['exists'] for r in results))

    def test_cache_invalidates_after_save_and_clear(self):
        self.assertEqual(excel.read_excel_summary(self.path)['stats']['total_invoices'], 0)
        self.assertTrue(excel.save_invoices_to_excel([self.invoice], self.path)['success'])
        saved = excel.read_excel_summary(self.path)
        self.assertEqual(saved['stats']['total_amount'], 1100)
        self.assertEqual(saved['rows'][0]['so_hd'], '123')
        self.assertTrue(excel.clear_excel_data(self.path)['success'])
        self.assertEqual(excel.read_excel_summary(self.path)['stats']['total_invoices'], 0)

    def test_external_edit_invalidates_cache(self):
        excel.save_invoices_to_excel([self.invoice], self.path)
        excel.read_excel_summary(self.path)
        wb = openpyxl.load_workbook(self.path)
        wb['TongHopHoaDon'].cell(2, 17, 2200)
        wb.save(self.path)
        wb.close()
        self.assertEqual(excel.read_excel_summary(self.path)['stats']['total_amount'], 2200)

    def test_streaming_workbook_without_dimension_metadata(self):
        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet('TongHopHoaDon')
        ws.append([h[0] for h in excel.SUMMARY_HEADERS])
        ws.append([1, 'NCC', '2026-10-07', '1', 'C26TAA', '123', '',
                   '0123456789', 'NCC', '', '', '', '', '', 1000, 100, 1100])
        wb.save(self.path)
        wb.close()
        summary = excel.read_excel_summary(self.path)
        self.assertNotIn('error', summary)
        self.assertEqual(summary['stats']['total_amount'], 1100)

    def test_corrupt_file_is_preserved_and_error_not_cached(self):
        original = b'broken workbook that must never be overwritten'
        Path(self.path).write_bytes(original)
        with self.assertRaises(ValueError):
            excel.ensure_excel_file(self.path)
        self.assertFalse(excel.save_invoices_to_excel([self.invoice], self.path)['success'])
        self.assertEqual(Path(self.path).read_bytes(), original)
        self.assertIn('error', excel.read_excel_summary(self.path))
        excel.init_blank_excel_file(self.path)
        self.assertNotIn('error', excel.read_excel_summary(self.path))

    def test_skip_replace_and_summary_values(self):
        self.assertEqual(excel.save_invoices_to_excel([self.invoice], self.path)['added'], 1)
        self.assertEqual(excel.save_invoices_to_excel([self.invoice], self.path)['skipped'], 1)
        self.assertEqual(excel.save_invoices_to_excel([self.invoice], self.path, overwrite=True)['updated'], 1)
        summary = excel.read_excel_summary(self.path)
        self.assertEqual(summary['stats']['total_invoices'], 1)
        self.assertEqual(summary['stats']['total_vat'], 100)
        self.assertEqual(summary['pivot_by_month'][0]['month'], '2026-10')

    def test_api_startup_and_mutations_are_isolated(self):
        import app as application
        settings = str(Path(self.temp.name) / 'settings.json')
        with patch.object(application, 'current_excel_path', self.path), patch.object(application, 'SETTINGS_FILE', settings):
            client = application.app.test_client()
            html = client.get('/').get_data(as_text=True)
            self.assertIn('id="startup-screen"', html)
            import config
            self.assertIn(f'v{config.APP_VERSION}', html)
            self.assertNotIn('cdn.tailwindcss.com', html)
            self.assertEqual(client.get('/api/status').status_code, 200)
            self.assertEqual(client.get('/api/status').get_json()['metadata']['version'], config.APP_VERSION)
            self.assertEqual(client.get('/api/excel-data').get_json()['stats']['total_invoices'], 0)
            self.assertTrue(client.post('/api/save-to-excel', json={'invoices': [self.invoice]}).get_json()['success'])
            self.assertEqual(client.post('/api/clear-data', json={'confirmation': 'wrong'}).status_code, 400)
            self.assertEqual(client.get('/api/excel-data').get_json()['stats']['total_invoices'], 1)
            self.assertTrue(client.post('/api/clear-data', json={'confirmation': 'XÓA TOÀN BỘ DỮ LIỆU'}).get_json()['success'])
            for asset in ['/static/css/tailwind.css', '/static/css/startup.css', '/static/vendor/fontawesome/webfonts/fa-solid-900.woff2']:
                with client.get(asset) as response:
                    self.assertEqual(response.status_code, 200)

    def test_desktop_splash_embeds_all_assets(self):
        from desktop_app import startup_html
        html = startup_html()
        self.assertIn('data:image/png;base64,', html)
        self.assertNotIn('src="/static/', html)
        self.assertIn('startup-loading', html)


if __name__ == '__main__':
    unittest.main()
