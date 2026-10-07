import os
import openpyxl
from parser import parse_invoice_file, scan_folder_for_invoices
from excel_manager import save_invoices_to_excel, read_excel_summary, ensure_excel_file
from app import app

def run_tests():
    client = app.test_client()

    # 1. Test status API
    res = client.get('/api/status')
    assert res.status_code == 200, 'Status API failed'
    print('Test 1: Status API OK')

    # 2. Test Excel Data API
    res = client.get('/api/excel-data')
    assert res.status_code == 200, 'Excel Data API failed'
    data = res.get_json()
    assert 'pivot_by_supplier' in data, 'Missing pivot_by_supplier'
    assert 'pivot_by_month' in data, 'Missing pivot_by_month'
    print('Test 2: Excel Data API with Pivot OK')

    # 3. Test duplicate behavior and Sheet structure
    test_excel = 'data/test_duplicate_check.xlsx'
    if os.path.exists(test_excel):
        os.remove(test_excel)
    ensure_excel_file(test_excel)

    sample_xml = r'C:\Users\minht\Downloads\348826-T01-2026.xml'
    if os.path.exists(sample_xml):
        inv1 = parse_invoice_file(sample_xml)
        inv1['supplier_folder'] = 'NhaCungCap_Aeon'
        
        # Save once
        s1 = save_invoices_to_excel([inv1], test_excel, overwrite=False)
        assert s1['added'] == 1, 'First save should add 1'
        
        # Save duplicate with overwrite=False (Skip)
        s2 = save_invoices_to_excel([inv1], test_excel, overwrite=False)
        assert s2['skipped'] == 1 and s2['added'] == 0, 'Duplicate with skip failed'
        
        # Save duplicate with overwrite=True (Replace in-place)
        s3 = save_invoices_to_excel([inv1], test_excel, overwrite=True)
        assert s3['updated'] == 1 and s3['added'] == 0, 'Duplicate with replace failed'
        
        # Check total rows
        sum_data = read_excel_summary(test_excel)
        tot = sum_data['stats']['total_invoices']
        assert tot == 1, f'Total invoices should be 1, got {tot}'
        
        # Check sheets
        wb = openpyxl.load_workbook(test_excel)
        assert wb.sheetnames[0] == 'TongQuan', f'Sheet 1 should be TongQuan, got {wb.sheetnames[0]}'
        assert 'TongHopHoaDon' in wb.sheetnames, 'TongHopHoaDon missing'
        assert 'ChiTietHangHoa' in wb.sheetnames, 'ChiTietHangHoa missing'
        
        ws_tq = wb['TongQuan']
        assert ws_tq.auto_filter.ref is not None, 'AutoFilter reference on TongQuan should not be None'
        print(f'Test 3: Duplicate logic & Sheet 1 TongQuan with AutoFilter ({ws_tq.auto_filter.ref}) OK')
        wb.close()
        
        if os.path.exists(test_excel):
            os.remove(test_excel)

    # 4. Test Scan folder API
    res = client.post('/api/scan-folder', json={
        'folder_path': r'C:\Users\minht\Downloads',
        'recursive': False,
        'auto_save': False,
        'overwrite': True
    })
    assert res.status_code == 200, f'Scan folder failed: {res.get_json()}'
    scan_json = res.get_json()
    assert 'supplier_folders' in scan_json, 'Missing supplier_folders in scan response'
    print(f'Test 4: Scan folder API OK. Found {scan_json["total_files"]} files, {len(scan_json["supplier_folders"])} folders')

    # 5. Test Settings & Theme API (Material 3 Persistence)
    res = client.get('/api/settings')
    assert res.status_code == 200, 'GET /api/settings failed'
    settings_data = res.get_json()
    assert 'settings' in settings_data, 'Missing settings in response'

    # Đổi sang theme tím Material 3
    res = client.post('/api/settings', json={'theme': 'purple'})
    assert res.status_code == 200, 'POST /api/settings failed'
    updated = res.get_json()
    assert updated['settings']['theme'] == 'purple', 'Theme not updated to purple'

    # Kiểm tra status API phản ánh theme mới
    res_status = client.get('/api/status')
    assert res_status.get_json().get('theme') == 'purple', 'Status API did not reflect updated theme'

    # Reset về theme rose mặc định
    client.post('/api/settings', json={'theme': 'rose'})
    print('Test 5: Settings & Material 3 Theme API OK')

    # 6. Test In-App Update Checker API & SemVer comparator
    from app import compare_versions
    assert compare_versions('2.1.0', '2.0.0') == 1, '2.1.0 should be > 2.0.0'
    assert compare_versions('2.0.0', '2.0.0') == 0, '2.0.0 should be == 2.0.0'
    assert compare_versions('1.9.9', '2.0.0') == -1, '1.9.9 should be < 2.0.0'
    assert compare_versions('v2.0.1', '2.0.0') == 1, 'v2.0.1 should be > 2.0.0'

    res_update = client.get('/api/check-update')
    assert res_update.status_code == 200, f'Check update failed with {res_update.status_code}: {res_update.get_json()}'
    update_data = res_update.get_json()
    assert update_data.get('success') is True, 'Check update API returned success: False'
    assert 'current_version' in update_data, 'Missing current_version'
    assert 'latest_version' in update_data, 'Missing latest_version'
    assert 'has_update' in update_data, 'Missing has_update'
    print(f'Test 6: In-App Update Checker API OK. Current={update_data["current_version"]}, Latest={update_data["latest_version"]}, HasUpdate={update_data["has_update"]}')

    print('>>> ALL TESTS PASSED SUCCESSFULLY! <<<')

if __name__ == '__main__':
    run_tests()
