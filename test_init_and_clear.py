import os
import sys
import json
import unittest

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

from excel_manager import (
    init_blank_excel_file,
    clear_excel_data,
    read_excel_summary,
    save_invoices_to_excel
)
from app import app
import openpyxl

class TestInitAndClear(unittest.TestCase):
    def setUp(self):
        self.test_excel = os.path.abspath("data/test_init_and_clear.xlsx")
        if os.path.exists(self.test_excel):
            os.remove(self.test_excel)
        self.client = app.test_client()

    def tearDown(self):
        if os.path.exists(self.test_excel):
            try:
                os.remove(self.test_excel)
            except Exception:
                pass

    def test_init_blank_excel_file(self):
        print("\n--- 1. Kiểm tra khởi tạo file Excel mới trắng tinh ---")
        res = init_blank_excel_file(self.test_excel)
        self.assertTrue(res['success'], f"Lỗi khởi tạo: {res.get('error')}")
        self.assertTrue(os.path.exists(self.test_excel))

        wb = openpyxl.load_workbook(self.test_excel)
        self.assertIn("TongQuan", wb.sheetnames)
        self.assertIn("TongHopHoaDon", wb.sheetnames)
        self.assertIn("ChiTietHangHoa", wb.sheetnames)
        
        # Kiểm tra headers
        ws_sum = wb["TongHopHoaDon"]
        self.assertEqual(ws_sum.max_row, 1, "TongHopHoaDon chỉ nên có dòng 1 (header)")
        ws_det = wb["ChiTietHangHoa"]
        self.assertEqual(ws_det.max_row, 1, "ChiTietHangHoa chỉ nên có dòng 1 (header)")
        wb.close()

        summary = read_excel_summary(self.test_excel)
        self.assertEqual(summary['stats']['total_invoices'], 0)
        self.assertEqual(summary['stats']['total_amount'], 0.0)
        print("✓ Khởi tạo file Excel mới thành công với đầy đủ 3 Sheet chuẩn mực!")

    def test_clear_excel_data(self):
        print("\n--- 2. Kiểm tra xóa toàn bộ dữ liệu Excel ---")
        # Khởi tạo và ghi 1 hóa đơn thật từ SAMPLE_XML_PATH
        init_blank_excel_file(self.test_excel)
        import config
        from parser import parse_invoice_file
        parsed = parse_invoice_file(config.SAMPLE_XML_PATH)
        self.assertTrue(parsed['success'])
        save_invoices_to_excel([parsed], self.test_excel)
        
        # Kiểm tra trước khi xóa
        summary_before = read_excel_summary(self.test_excel)
        self.assertEqual(summary_before['stats']['total_invoices'], 1)
        self.assertGreater(summary_before['stats']['total_amount'], 0)

        # Thực hiện xóa
        clear_res = clear_excel_data(self.test_excel)
        self.assertTrue(clear_res['success'], f"Lỗi xóa data: {clear_res.get('error')}")

        # Kiểm tra sau khi xóa
        summary_after = read_excel_summary(self.test_excel)
        self.assertEqual(summary_after['stats']['total_invoices'], 0)
        self.assertEqual(summary_after['stats']['total_amount'], 0.0)
        self.assertEqual(len(summary_after['rows']), 0)

        # Kiểm tra tiêu đề vẫn nguyên vẹn
        wb = openpyxl.load_workbook(self.test_excel)
        self.assertEqual(wb["TongHopHoaDon"].max_row, 1)
        self.assertEqual(wb["ChiTietHangHoa"].max_row, 1)
        wb.close()
        print("✓ Xóa dữ liệu thành công! Hóa đơn về 0, headers và 3 Sheet vẫn bảo toàn 100%!")

    def test_api_clear_data_strict_constraint(self):
        print("\n--- 3. Kiểm tra API /api/clear-data ràng buộc chuỗi nghiêm ngặt ---")
        init_blank_excel_file(self.test_excel)
        
        # Test chuỗi sai -> Bị chặn
        res_fail_1 = self.client.post('/api/clear-data', json={'confirmation': 'xoa het'})
        self.assertEqual(res_fail_1.status_code, 400)
        data_fail_1 = json.loads(res_fail_1.data)
        self.assertFalse(data_fail_1['success'])
        self.assertIn('Chuỗi xác nhận không chính xác', data_fail_1['error'])

        res_fail_2 = self.client.post('/api/clear-data', json={'confirmation': 'xóa toàn bộ dữ liệu'}) # Chữ thường
        self.assertEqual(res_fail_2.status_code, 400)

        # Test chuỗi đúng -> Thành công
        res_ok = self.client.post('/api/clear-data', json={'confirmation': 'XÓA TOÀN BỘ DỮ LIỆU'})
        self.assertEqual(res_ok.status_code, 200)
        data_ok = json.loads(res_ok.data)
        self.assertTrue(data_ok['success'])
        self.assertEqual(data_ok['stats']['total_invoices'], 0)
        print("✓ API /api/clear-data chặn mọi chuỗi sai và chỉ kích hoạt khi gõ đúng 100% 'XÓA TOÀN BỘ DỮ LIỆU'!")

    def test_api_init_excel_file(self):
        print("\n--- 4. Kiểm tra API /api/init-excel-file ---")
        custom_target = os.path.abspath("data/custom_init_test.xlsx")
        if os.path.exists(custom_target):
            os.remove(custom_target)

        res = self.client.post('/api/init-excel-file', json={'excel_path': custom_target})
        self.assertEqual(res.status_code, 200)
        data = json.loads(res.data)
        self.assertTrue(data['success'])
        self.assertTrue(os.path.exists(custom_target))
        self.assertEqual(data['stats']['total_invoices'], 0)

        # Cleanup
        if os.path.exists(custom_target):
            os.remove(custom_target)
        print("✓ API /api/init-excel-file tạo file mới và đồng bộ trạng thái server thành công!")

if __name__ == '__main__':
    unittest.main()
