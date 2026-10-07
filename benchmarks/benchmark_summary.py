import json, subprocess, tempfile, time
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
import excel_manager as current
old = {'__name__': 'baseline_excel'}
exec(compile(subprocess.check_output(['git', 'show', '5c90279:excel_manager.py']).decode('utf-8'), 'baseline_excel.py', 'exec'), old)
results=[]
with tempfile.TemporaryDirectory() as tmp:
    for count in (1000, 10000):
        p=str(Path(tmp)/f'benchmark-{count}.xlsx')
        wb=openpyxl.Workbook(write_only=True)
        ws=wb.create_sheet('TongHopHoaDon')
        ws.append([h[0] for h in current.SUMMARY_HEADERS])
        for i in range(count):
            ws.append([i+1,'NCC','2026-10-07','1','C26TAA',str(i+1),'','0123456789','NCC','Address','9876543210','Buyer','Address','TM',1000,100,1100,'',5,'XML','sample.xml','2026-10-07'])
        detail=wb.create_sheet('ChiTietHangHoa')
        detail.append([h[0] for h in current.DETAIL_HEADERS])
        for i in range(count*5):
            detail.append([i+1,'NCC','C26TAA',str(i//5+1),'2026-10-07','0123456789','NCC','2026-10-07','SP','Item','Unit',1,200,200,'10%',20,220])
        wb.save(p);wb.close()
        t=time.perf_counter();before=old['read_excel_summary'](p);baseline=time.perf_counter()-t
        current._summary_cache.clear()
        t=time.perf_counter();after=current.read_excel_summary(p);cold=time.perf_counter()-t
        t=time.perf_counter();cached=current.read_excel_summary(p);warm=time.perf_counter()-t
        assert before['rows']==after['rows'] and before['stats']==after['stats'], (before.get('error'),after.get('error'))
        results.append({'invoices':count,'detail_rows':count*5,'baseline_seconds':round(baseline,4),'optimized_cold_seconds':round(cold,4),'cached_seconds':round(warm,4)})
        print(json.dumps(results[-1]), flush=True)
Path('build').mkdir(exist_ok=True)
Path('build/benchmark-results.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
